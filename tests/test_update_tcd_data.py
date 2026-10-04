from pathlib import Path
import json
import shutil
import sqlite3

import pandas as pd
import pytest
from openpyxl import Workbook

from scripts.update_tcd_data import (
    NIGHTS_BIN_ORDER,
    build_tcd_sqlite,
    extract_t06_rows,
    read_t06_unit,
    validate_tcd_rows,
)
from scripts import update_tcd_data as updater


def annual_book_with_quarter_detail() -> Workbook:
    book = Workbook()
    sheet = book.active
    sheet.title = "T06"
    sheet["I2"] = "（千泊）"
    sheet["A65"] = "2025年"
    sheet["A68"] = "宿泊数"
    for index, label in enumerate(NIGHTS_BIN_ORDER, start=69):
        sheet.cell(index, 1, label)
        sheet.cell(index, 2, 100 + index)
        sheet.cell(index, 5, 10 + index)
    sheet["A85"] = "2025年1～3月期"
    sheet["A148"] = "宿泊数"
    for index, label in enumerate(NIGHTS_BIN_ORDER, start=149):
        sheet.cell(index, 1, label)
        sheet.cell(index, 2, 200 + index)
        sheet.cell(index, 5, 20 + index)
    return book


def test_annual_workbook_extracts_annual_section_without_quarter_double_count() -> None:
    book = annual_book_with_quarter_detail()
    assert read_t06_unit(book["T06"]) == "千泊"
    rows = extract_t06_rows(
        book, "https://example.test/annual.xlsx",
        "2025年確報", "sha", ("annual", "2025", "2025年"), "確報"
    )
    assert len(rows) == 16
    assert set(rows["period_key"]) == {"2025"}
    assert rows.loc[(rows.segment == "domestic_total") & (rows.nights_bin == "1泊"), "value"].item() == 169


def test_unknown_t06_unit_is_rejected() -> None:
    book = annual_book_with_quarter_detail()
    book["T06"]["I2"] = "（百万円）"
    with pytest.raises(ValueError, match="unit is not supported"):
        read_t06_unit(book["T06"])


def test_duplicate_period_segment_bin_rejected_before_existing_db_is_replaced(tmp_path: Path) -> None:
    db = tmp_path / "market_stats.sqlite"
    original = pd.DataFrame([{
        "period_type":"annual", "period_key":"2024", "period_label":"2024年",
        "release_type":"確報", "segment":"domestic_total", "nights_bin":"1泊",
        "value":1.0, "source_url":"old", "source_title":"old", "source_sha256":"old"
    }])
    build_tcd_sqlite(original, db)
    duplicate = pd.concat([original, original.assign(value=2.0)], ignore_index=True)
    with pytest.raises(ValueError, match="Duplicate TCD grain"):
        build_tcd_sqlite(duplicate, db)
    with sqlite3.connect(db) as conn:
        assert conn.execute("select value from tcd_stay_nights").fetchall() == [(1.0,)]


def test_tcd_rejects_missing_numeric_values_and_inconsistent_period_keys() -> None:
    row = {
        "period_type":"annual", "period_key":"2025", "period_label":"2025年",
        "release_type":"確報", "segment":"domestic_total", "nights_bin":"1泊",
        "value":1.0, "source_url":"x", "source_title":"x", "source_sha256":"x"
    }
    with pytest.raises(ValueError, match="Missing TCD value"):
        validate_tcd_rows(pd.DataFrame([row | {"value":float("nan")}]))
    with pytest.raises(ValueError, match="period key"):
        validate_tcd_rows(pd.DataFrame([row | {"period_key":"2025Q1"}]))


def test_reprocessing_annual_source_is_idempotent_and_keeps_only_annual_grain(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "annual.xlsx"
    book = annual_book_with_quarter_detail()
    book.create_sheet("表題")["A1"] = "2025年旅行・観光消費動向調査（確報）"
    book.save(source)
    data = tmp_path / "data"
    data.mkdir()
    monkeypatch.setattr(updater, "DATA_DIR", data)
    monkeypatch.setattr(updater, "SQLITE_PATH", data / "market_stats.sqlite")
    monkeypatch.setattr(updater, "META_TCD_PATH", data / "meta_tcd.json")
    monkeypatch.setattr(updater, "fetch_source_page", lambda url: "")
    monkeypatch.setattr(updater, "extract_target_excel_links", lambda html, url: [{
        "url": "https://example.test/annual.xlsx", "link_text": "2025年（確報）集計表"
    }])
    monkeypatch.setattr(updater, "download_file", lambda url, dst: shutil.copyfile(source, dst))

    assert updater.main() == 0
    with sqlite3.connect(data / "market_stats.sqlite") as conn:
        first = conn.execute("select period_key,segment,nights_bin,value from tcd_stay_nights order by 1,2,3").fetchall()
    assert len(first) == 16
    assert {row[0] for row in first} == {"2025"}
    assert updater.main() == 0
    with sqlite3.connect(data / "market_stats.sqlite") as conn:
        second = conn.execute("select period_key,segment,nights_bin,value from tcd_stay_nights order by 1,2,3").fetchall()
    assert first == second
    assert json.loads((data / "meta_tcd.json").read_text())["extractor_version"] == updater.EXTRACTOR_VERSION
    assert json.loads((data / "meta_tcd.json").read_text())["processed_files"][0]["source_unit"] == "千泊"
    source.write_bytes(b"broken workbook")
    with pytest.raises(RuntimeError, match="Previously published TCD source could not be opened"):
        updater.main()
    with sqlite3.connect(data / "market_stats.sqlite") as conn:
        preserved = conn.execute("select period_key,segment,nights_bin,value from tcd_stay_nights order by 1,2,3").fetchall()
    assert preserved == first
