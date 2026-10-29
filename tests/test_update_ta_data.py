import json
import shutil
import sqlite3
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scripts import update_ta_data as updater


def write_ta_workbook(path: Path, companies: list[str]) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    for column, label in [
        (3, "海外旅行"),
        (8, "外国人旅行"),
        (13, "国内旅行"),
        (18, "合計"),
    ]:
        sheet.cell(4, column, label)
        sheet.cell(5, column, "2024年度")
        sheet.cell(6, column, "取扱額（千円）")
    for row, company in enumerate(companies, start=7):
        sheet.cell(row, 2, company)
        for column, amount in [(3, 10), (8, 20), (13, 30), (18, 60)]:
            sheet.cell(row, column, amount)
    book.save(path)
    book.close()


@pytest.mark.parametrize(
    "summary", ["合計", "合 計", "　合　計　", "小計", "小 計", "\t小\u00a0計\n"]
)
def test_parser_excludes_summary_company_rows(tmp_path: Path, summary: str) -> None:
    source = tmp_path / "companies.xlsx"
    write_ta_workbook(source, ["JTB（6社計）", summary])

    rows = updater.parse_ta_excel_rows(source, "2024", "total")

    assert rows == [
        {
            "fiscal_year": "2024",
            "period": "total",
            "company": "JTB（6社計）",
            "segment": "overseas",
            "amount": 10.0,
        },
        {
            "fiscal_year": "2024",
            "period": "total",
            "company": "JTB（6社計）",
            "segment": "foreign",
            "amount": 20.0,
        },
        {
            "fiscal_year": "2024",
            "period": "total",
            "company": "JTB（6社計）",
            "segment": "domestic",
            "amount": 30.0,
        },
        {
            "fiscal_year": "2024",
            "period": "total",
            "company": "JTB（6社計）",
            "segment": "total",
            "amount": 60.0,
        },
    ]


@pytest.fixture
def update_paths(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> tuple[Path, Path]:
    source = tmp_path / "companies.xlsx"
    write_ta_workbook(source, ["JTB（6社計）", "合 計"])
    data_dir = tmp_path / "data"
    sqlite_path = data_dir / "market_stats.sqlite"
    meta_path = data_dir / "meta_ta.json"
    monkeypatch.setattr(updater, "DATA_DIR", data_dir)
    monkeypatch.setattr(updater, "SQLITE_PATH", sqlite_path)
    monkeypatch.setattr(updater, "META_PATH", meta_path)

    fiscal_url = "https://example.test/ryokogyotoriatsukaigaku_r6.html"
    excel_url = "https://example.test/companies.xlsx"
    pages = {
        updater.TA_SOURCE_INDEX_URL: f'<a href="{fiscal_url}">2024年度</a>',
        fiscal_url: f'<title>2024年度</title><a href="{excel_url}">2024年度総計 各社別内訳</a>',
    }
    monkeypatch.setattr(updater, "fetch_html", pages.__getitem__)

    def download_fixture(url: str, destination: Path) -> None:
        if url != excel_url:
            raise ValueError(f"Unexpected fixture URL: {url}")
        shutil.copyfile(source, destination)

    monkeypatch.setattr(updater, "download_file", download_fixture)
    return sqlite_path, meta_path


def read_stored_rows(sqlite_path: Path) -> list[tuple]:
    with sqlite3.connect(sqlite_path) as conn:
        return conn.execute(
            "SELECT company, segment, amount FROM ta_company_amounts ORDER BY segment"
        ).fetchall()


@pytest.mark.parametrize("old_version", [None, 0, 1, 2, 3])
def test_unchanged_source_reprocesses_with_missing_or_old_extractor_version(
    update_paths: tuple[Path, Path], old_version: int | None
) -> None:
    sqlite_path, meta_path = update_paths
    assert updater.main() == 0
    old_meta = json.loads(meta_path.read_text(encoding="utf-8"))
    if old_version is None:
        old_meta.pop("extractor_version", None)
    else:
        old_meta["extractor_version"] = old_version
    meta_path.write_text(json.dumps(old_meta), encoding="utf-8")
    with sqlite3.connect(sqlite_path) as conn:
        conn.execute(
            "INSERT INTO ta_company_amounts VALUES ('2024', 'total', '合 計', 'total', 60)"
        )

    assert updater.main() == 0

    assert read_stored_rows(sqlite_path) == [
        ("JTB（6社計）", "domestic", 30.0),
        ("JTB（6社計）", "foreign", 20.0),
        ("JTB（6社計）", "overseas", 10.0),
        ("JTB（6社計）", "total", 60.0),
    ]
    new_meta = json.loads(meta_path.read_text(encoding="utf-8"))
    assert new_meta["signature"] == old_meta["signature"]
    assert new_meta["extractor_version"] == updater.EXTRACTOR_VERSION
    assert new_meta["row_count"] == 4


@pytest.mark.parametrize(
    "table_state", ["missing_database", "missing_table", "empty_table"]
)
def test_unchanged_source_rebuilds_an_unusable_existing_table(
    update_paths: tuple[Path, Path], table_state: str
) -> None:
    sqlite_path, meta_path = update_paths
    assert updater.main() == 0
    old_meta = json.loads(meta_path.read_text(encoding="utf-8"))
    old_meta["extractor_version"] = updater.EXTRACTOR_VERSION
    meta_path.write_text(json.dumps(old_meta), encoding="utf-8")
    if table_state == "missing_database":
        sqlite_path.unlink()
    else:
        with sqlite3.connect(sqlite_path) as conn:
            if table_state == "missing_table":
                conn.execute("DROP TABLE ta_company_amounts")
            else:
                conn.execute("DELETE FROM ta_company_amounts")

    assert updater.main() == 0

    assert sqlite_path.exists()
    assert read_stored_rows(sqlite_path) == [
        ("JTB（6社計）", "domestic", 30.0),
        ("JTB（6社計）", "foreign", 20.0),
        ("JTB（6社計）", "overseas", 10.0),
        ("JTB（6社計）", "total", 60.0),
    ]
    new_meta = json.loads(meta_path.read_text(encoding="utf-8"))
    assert new_meta["extractor_version"] == updater.EXTRACTOR_VERSION


def test_current_extractor_and_usable_table_skip_unchanged_source(
    update_paths: tuple[Path, Path],
) -> None:
    sqlite_path, meta_path = update_paths
    assert updater.main() == 0
    old_meta = json.loads(meta_path.read_text(encoding="utf-8"))
    old_meta["extractor_version"] = updater.EXTRACTOR_VERSION
    old_meta["last_checked_at"] = "previous-successful-parse"
    meta_path.write_text(json.dumps(old_meta), encoding="utf-8")
    database_before = sqlite_path.read_bytes()
    metadata_before = meta_path.read_bytes()

    assert updater.main() == 0

    assert sqlite_path.read_bytes() == database_before
    assert meta_path.read_bytes() == metadata_before


def test_successful_update_persists_extractor_version(
    update_paths: tuple[Path, Path],
) -> None:
    _, meta_path = update_paths

    assert updater.main() == 0

    new_meta = json.loads(meta_path.read_text(encoding="utf-8"))
    assert new_meta["extractor_version"] == updater.EXTRACTOR_VERSION


@pytest.fixture
def published_two_periods(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> tuple[Path, Path, dict[str, str], dict[str, Path]]:
    annual = tmp_path / "annual.xlsx"
    monthly = tmp_path / "monthly.xlsx"
    write_ta_workbook(annual, ["JTB（6社計）"])
    write_ta_workbook(monthly, ["会社B"])
    data_dir = tmp_path / "data"
    sqlite_path = data_dir / "market_stats.sqlite"
    meta_path = data_dir / "meta_ta.json"
    monkeypatch.setattr(updater, "DATA_DIR", data_dir)
    monkeypatch.setattr(updater, "SQLITE_PATH", sqlite_path)
    monkeypatch.setattr(updater, "META_PATH", meta_path)
    fiscal_url = "https://example.test/ryokogyotoriatsukaigaku_r6.html"
    source_paths = {
        "https://example.test/annual.xlsx": annual,
        "https://example.test/monthly.xlsx": monthly,
    }
    pages = {
        updater.TA_SOURCE_INDEX_URL: f'<a href="{fiscal_url}">2024年度</a>',
        fiscal_url: (
            "<title>2024年度</title>"
            '<a href="https://example.test/annual.xlsx">2024年度総計 各社別内訳</a>'
            '<a href="https://example.test/monthly.xlsx">2024年4月 各社別内訳</a>'
        ),
    }
    monkeypatch.setattr(updater, "fetch_html", pages.__getitem__)
    monkeypatch.setattr(
        updater,
        "download_file",
        lambda url, dst: shutil.copyfile(source_paths[url], dst),
    )
    assert updater.main() == 0
    return sqlite_path, meta_path, pages, source_paths


@pytest.mark.parametrize("failure", ["invalid_workbook", "no_company_rows"])
@pytest.mark.parametrize("replacement_url", [False, True])
def test_partial_parse_failure_preserves_database_and_metadata(
    published_two_periods: tuple[Path, Path, dict[str, str], dict[str, Path]],
    failure: str,
    replacement_url: bool,
) -> None:
    sqlite_path, meta_path, pages, sources = published_two_periods
    monthly_url = "https://example.test/monthly.xlsx"
    if replacement_url:
        new_url = "https://example.test/monthly_replacement.xlsx"
        sources[new_url] = sources.pop(monthly_url)
        fiscal_url = "https://example.test/ryokogyotoriatsukaigaku_r6.html"
        pages[fiscal_url] = pages[fiscal_url].replace(monthly_url, new_url)
        monthly_url = new_url
    if failure == "invalid_workbook":
        sources[monthly_url].write_bytes(b"invalid xlsx workbook")
    else:
        write_ta_workbook(sources[monthly_url], [])
    database_before = sqlite_path.read_bytes()
    metadata_before = meta_path.read_bytes()

    with pytest.raises(RuntimeError, match="TA source"):
        updater.main()

    assert sqlite_path.read_bytes() == database_before
    assert meta_path.read_bytes() == metadata_before


@pytest.mark.parametrize(
    "failure", ["unsupported_xls", "unknown_fiscal_year", "unknown_period"]
)
def test_unclassified_sources_preserve_database_and_metadata(
    published_two_periods: tuple[Path, Path, dict[str, str], dict[str, Path]],
    failure: str,
) -> None:
    sqlite_path, meta_path, pages, sources = published_two_periods
    fiscal_url = "https://example.test/ryokogyotoriatsukaigaku_r6.html"
    if failure == "unsupported_xls":
        xlsx_url = "https://example.test/monthly.xlsx"
        xls_url = "https://example.test/monthly.xls"
        sources[xls_url] = sources.pop(xlsx_url)
        pages[fiscal_url] = pages[fiscal_url].replace(xlsx_url, xls_url)
    elif failure == "unknown_fiscal_year":
        for url in pages:
            pages[url] = pages[url].replace("2024年度", "当年度")
    else:
        pages[fiscal_url] = pages[fiscal_url].replace("2024年4月", "期間不明")
    database_before = sqlite_path.read_bytes()
    metadata_before = meta_path.read_bytes()

    with pytest.raises(RuntimeError, match="TA source"):
        updater.main()

    assert sqlite_path.read_bytes() == database_before
    assert meta_path.read_bytes() == metadata_before


def test_metadata_write_failure_preserves_pair_and_retry_is_idempotent(
    published_two_periods: tuple[Path, Path, dict[str, str], dict[str, Path]],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    sqlite_path, meta_path, _, sources = published_two_periods
    source = sources["https://example.test/monthly.xlsx"]
    write_ta_workbook(source, ["会社C"])
    source_before = source.read_bytes()
    database_before = sqlite_path.read_bytes()
    metadata_before = meta_path.read_bytes()
    write_text = Path.write_text

    def fail_metadata(path: Path, *args, **kwargs):
        if path.name == "meta_ta.json":
            raise OSError("injected metadata write failure")
        return write_text(path, *args, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(Path, "write_text", fail_metadata)
        with pytest.raises(OSError, match="injected metadata"):
            updater.main()

    assert sqlite_path.read_bytes() == database_before
    assert meta_path.read_bytes() == metadata_before
    assert source.read_bytes() == source_before
    assert updater.main() == 0
    with sqlite3.connect(sqlite_path) as conn:
        assert conn.execute("SELECT DISTINCT company FROM ta_company_amounts ORDER BY 1").fetchall() == [
            ("JTB（6社計）",), ("会社C",)
        ]
        assert conn.execute("SELECT count(*) FROM ta_company_amounts").fetchone() == (8,)
    successful = sqlite_path.read_bytes(), meta_path.read_bytes()
    assert updater.main() == 0
    assert (sqlite_path.read_bytes(), meta_path.read_bytes()) == successful


def write_annual_ta_workbook(path: Path, *, previous_first: bool = False) -> None:
    # Synthetic fixture matching the layout of MLIT 002006785.xlsx (2025 total).
    book = Workbook()
    sheet = book.active
    sheet["A5"] = "会　社　名"
    for col, label, amount in [(2, "海外旅行", 10), (5, "外国人旅行", 20),
                                (8, "国内旅行", 30), (11, "合計", 60)]:
        sheet.cell(4, col, label)
        for offset, year in enumerate(([2024, 2025] if previous_first else [2025, 2024])):
            sheet.cell(5, col + offset, f"{year}年度取扱額")
            sheet.cell(6, col + offset, "（千円）")
            sheet.cell(7, col + offset, amount if year == 2025 else 999)
            sheet.cell(8, col + offset, amount)
        sheet.cell(5, col + 2, "2024年度")
        sheet.cell(6, col + 2, "比％")
    sheet["A7"] = "会社A"
    sheet["A8"] = "合　計"
    book.save(path)
    book.close()


@pytest.mark.parametrize("previous_first", [False, True])
def test_annual_layout_selects_company_and_requested_year(tmp_path, previous_first):
    source = tmp_path / "annual.xlsx"
    write_annual_ta_workbook(source, previous_first=previous_first)
    rows = updater.parse_ta_excel_rows(source, "2025", "total")
    assert [(r["company"], r["segment"], r["amount"]) for r in rows] == [
        ("会社A", "overseas", 10.0), ("会社A", "foreign", 20.0),
        ("会社A", "domestic", 30.0), ("会社A", "total", 60.0),
    ]
    assert all(r["fiscal_year"] == "2025" and r["period"] == "total" for r in rows)


def test_annual_layout_rejects_missing_requested_year(tmp_path):
    source = tmp_path / "annual.xlsx"
    write_annual_ta_workbook(source)
    with pytest.raises(ValueError, match="Amount columns"):
        updater.parse_ta_excel_rows(source, "2026", "total")


def test_annual_layout_does_not_borrow_next_segment_column(tmp_path):
    source = tmp_path / "annual.xlsx"
    write_annual_ta_workbook(source)
    book = load_workbook(source)
    book.active["B5"] = "unknown"
    book.save(source)
    book.close()
    with pytest.raises(ValueError, match="Amount columns"):
        updater.parse_ta_excel_rows(source, "2025", "total")


@pytest.mark.parametrize("period,previous_headers,previous_units", [
    ("2024-07", ["2023年同月", "2019年同月"], ["取扱額（千円）", "取扱額（千円）"]),
    ("2023-02", ["2019年同月"], ["取扱額（千円）"]),
    ("2022-04", ["2021年同月取", "2019年同月取"], ["扱額（千円）", "扱額（千円）"]),
    ("2021-04", ["前年同月取", "前々年同月取"], ["扱額（千円）", "扱額（千円）"]),
    ("total", ["2022年度取扱額", "2019年度取扱額"], ["（千円）", "（千円）"]),
])
def test_legacy_current_amount_in_row_five_has_priority(
    tmp_path, period, previous_headers, previous_units
):
    # Synthetic values, verbatim header shapes from MLIT 2021-2024 workbooks.
    source = tmp_path / "legacy.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["B5"] = "会　社　名"
    for start, label, amount in [(3, "海外旅行", 13208650), (8, "外国人旅行", 20),
                                  (13, "国内旅行", 30), (18, "合計", 60)]:
        sheet.cell(4, start, label)
        sheet.cell(5, start, "取　扱　額")
        sheet.cell(6, start, "（千円）")
        sheet.cell(7, start, amount)
        for offset, (header, unit) in enumerate(zip(previous_headers, previous_units), start=1):
            sheet.cell(5, start + offset, header)
            sheet.cell(6, start + offset, unit)
            sheet.cell(7, start + offset, 44771108)
    sheet["B7"] = "JTB"
    book.save(source)
    book.close()
    rows = updater.parse_ta_excel_rows(source, "2023", period)
    assert [(r["company"], r["segment"], r["amount"]) for r in rows] == [
        ("JTB", "overseas", 13208650.0), ("JTB", "foreign", 20.0),
        ("JTB", "domestic", 30.0), ("JTB", "total", 60.0),
    ]


@pytest.mark.parametrize("corrected,cancelled", [
    ("0\n", "343,304"), ("343,304\n", "0"),
    ("9,458\n", "13,604"), ("1,283,097\n", "1,278,950"),
    ("1,388,477\n", "1,388,476"), ("19,978\n", "14,376"),
    ("1,447,289\n", "1,452,890"), ("1,606,944\n", "1,606,943"),
    ("23,651\n", "5,194"), ("1,691,501\n", "1,709,958"),
    ("14,429\n", "0"), ("1,136,356\n", "1,150,785"),
])
def test_corrected_rich_text_amount_excludes_cancelled_run(corrected, cancelled):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    value = CellRichText(corrected, TextBlock(InlineFont(strike=True), cancelled))
    assert updater.to_float(value) == float(corrected.strip().replace(",", ""))


@pytest.mark.parametrize("storage", ["inline", "shared"])
def test_workbook_preserves_strike_correction_including_zero(tmp_path, storage):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from zipfile import ZipFile, ZIP_DEFLATED
    import xml.etree.ElementTree as ET

    source = tmp_path / "corrected.xlsx"
    write_ta_workbook(source, [])
    book = load_workbook(source)
    sheet = book.active
    sheet["B5"] = "会　社　名"
    for column in (3, 8, 13, 18):
        sheet.cell(5, column, "取　扱　額")
        sheet.cell(6, column, "（千円）")
    sheet["B28"] = "（株）IACEトラベル"
    sheet["C28"] = 100
    sheet["H28"] = CellRichText("0\n", TextBlock(InlineFont(strike=True), "343,304"))
    sheet["M28"] = CellRichText("343,304", "\n", TextBlock(InlineFont(strike=True), "0"))
    sheet["R28"] = 343404
    book.save(source)
    book.close()
    if storage == "shared":
        # Real MLIT corrections use sharedStrings runs, rather than the inline
        # rich text produced by openpyxl. Keep run properties intact in the fixture.
        ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        with ZipFile(source) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        root = ET.fromstring(files["xl/worksheets/sheet1.xml"])
        shared = ET.Element(f"{{{ns}}}sst", count="2", uniqueCount="2")
        index = 0
        for cell in root.iter(f"{{{ns}}}c"):
            if cell.get("r") not in {"H28", "M28"}:
                continue
            inline = cell.find(f"{{{ns}}}is")
            inline.tag = f"{{{ns}}}si"
            cell.remove(inline)
            shared.append(inline)
            cell.set("t", "s")
            ET.SubElement(cell, f"{{{ns}}}v").text = str(index)
            index += 1
        files["xl/worksheets/sheet1.xml"] = ET.tostring(root)
        files["xl/sharedStrings.xml"] = ET.tostring(shared)
        ct_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
        ct = ET.fromstring(files["[Content_Types].xml"])
        ET.SubElement(ct, f"{{{ct_ns}}}Override", PartName="/xl/sharedStrings.xml",
                      ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml")
        files["[Content_Types].xml"] = ET.tostring(ct)
        with ZipFile(source, "w", ZIP_DEFLATED) as archive:
            for name, data in files.items():
                archive.writestr(name, data)
    rows = updater.parse_ta_excel_rows(source, "2023", "2024-03")
    amounts = {r["segment"]: r["amount"] for r in rows}
    assert amounts == {"overseas": 100, "foreign": 0, "domestic": 343304, "total": 343404}


def test_unknown_amount_stops_and_preserves_database_and_metadata(published_two_periods):
    sqlite_path, meta_path, _, sources = published_two_periods
    before = sqlite_path.read_bytes(), meta_path.read_bytes()
    source = sources["https://example.test/monthly.xlsx"]
    book = load_workbook(source)
    book.active["H7"] = "23,000\n24,000"
    book.save(source)
    book.close()
    with pytest.raises(RuntimeError, match="TA source parse failed.*H7"):
        updater.main()
    assert (sqlite_path.read_bytes(), meta_path.read_bytes()) == before


@pytest.mark.parametrize("value", [None, "", "-", "－", "…"])
def test_explicit_missing_amounts_remain_missing(value):
    assert updater.to_float(value) is None


@pytest.mark.parametrize("value", ["NaN", "inf", float("nan"), float("inf"), "1,23", True])
def test_unknown_numeric_values_raise_instead_of_becoming_missing(value):
    with pytest.raises(ValueError):
        updater.to_float(value)


def test_unresolved_rich_text_with_two_active_numbers_stops():
    from openpyxl.cell.rich_text import CellRichText
    with pytest.raises(ValueError, match="Unrecognised amount"):
        updater.to_float(CellRichText("1,000\n", "2,000"))


def test_cancelled_amount_without_replacement_stops():
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    with pytest.raises(ValueError, match="no uncancelled value"):
        updater.to_float(CellRichText(TextBlock(InlineFont(strike=True), "1,000")))
